import io
import os
import pytest
from pathlib import Path
from PIL import Image, ImageDraw
import pymupdf
from pptx import Presentation
from pptx.util import Inches

from app.core.config import settings
from app.models.material import Material, MaterialType, ProcessingStatus
from app.models.knowledge_base import ModalityType
from app.db.repository import repository
from app.services.parsers.visual_extractor import visual_extractor
from app.services.ingestion import ingestion_service
from app.services.vision_service import vision_service
from fastapi.testclient import TestClient
from app.main import app

client = TestClient(app)


def _create_synthetic_diagram_png(width=400, height=300, label="Glycolysis Pathway") -> bytes:
    """Generates an in-memory PNG diagram with labeled shapes."""
    img = Image.new("RGB", (width, height), color=(240, 245, 255))
    draw = ImageDraw.Draw(img)
    draw.rectangle([(20, 20), (width - 20, height - 20)], outline=(30, 64, 175), width=3)
    draw.ellipse([(50, 50), (180, 150)], fill=(191, 219, 254), outline=(37, 99, 235), width=2)
    draw.rectangle([(220, 50), (360, 150)], fill=(187, 247, 208), outline=(22, 163, 74), width=2)
    draw.line([(180, 100), (220, 100)], fill=(15, 23, 42), width=3)
    draw.text((60, 90), "Glucose", fill=(15, 23, 42))
    draw.text((250, 90), "Pyruvate", fill=(15, 23, 42))
    draw.text((40, 220), label, fill=(30, 64, 175))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


@pytest.fixture
def sample_pdf_with_diagram(tmp_path) -> Path:
    """Creates a sample PDF with textbook text, embedded diagram, and figure caption."""
    pdf_path = tmp_path / "biochemistry_chapter4.pdf"
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)  # A4

    text = (
        "Chapter 4: Cellular Respiration and Glycolysis\n\n"
        "Glycolysis is the sequence of ten chemical reactions that converts glucose into pyruvate.\n"
        "During this pathway, two molecules of ATP and two molecules of NADH are produced.\n\n"
        "Figure 3.2: Glycolysis pathway showing enzymatic steps from glucose to pyruvate.\n\n"
        "The reaction sequence is tightly regulated by phosphofructokinase-1."
    )
    page.insert_text((50, 50), text, fontsize=12)

    # Insert synthetic diagram image into PDF page
    diagram_bytes = _create_synthetic_diagram_png(400, 300, "Figure 3.2: Glycolysis")
    img_rect = pymupdf.Rect(50, 250, 450, 550)
    page.insert_image(img_rect, stream=diagram_bytes)

    doc.save(str(pdf_path))
    doc.close()
    return pdf_path


@pytest.fixture
def sample_pptx_with_diagram(tmp_path) -> Path:
    """Creates a sample PPTX slide deck with a diagram picture."""
    pptx_path = tmp_path / "bio_lecture4.pptx"
    prs = Presentation()

    # Slide 1: Title
    slide1 = prs.slides.add_slide(prs.slide_layouts[0])
    slide1.shapes.title.text = "Lecture 4: Bioenergetics"
    slide1.placeholders[1].text = "Dr. Jane Doe - Principles of Biology"

    # Slide 2: Mitochondrial Chemiosmosis with embedded diagram picture
    slide2 = prs.slides.add_slide(prs.slide_layouts[5])  # Title only
    slide2.shapes.title.text = "Mitochondrial Chemiosmosis"

    diag_bytes = _create_synthetic_diagram_png(400, 300, "Mitochondrial Chemiosmosis")
    diag_stream = io.BytesIO(diag_bytes)
    slide2.shapes.add_picture(diag_stream, Inches(1.5), Inches(2.0), width=Inches(6.0))

    prs.save(str(pptx_path))
    return pptx_path


@pytest.mark.asyncio
async def test_vision_service_analysis():
    """Verify VisionService produces structured 5-part visual descriptions."""
    test_image = _create_synthetic_diagram_png()
    analysis = await vision_service.analyze_image(
        image_bytes=test_image,
        caption="Figure 3.2: Glycolysis pathway",
        context_text="Cellular respiration converting glucose to pyruvate net 2 ATP",
        source_info="Biochemistry Textbook p.42",
    )

    assert analysis["success"] is True
    assert "visual_type" in analysis
    assert analysis["visual_type"] in ["diagram", "chart", "figure", "table", "illustration"]
    assert "content" in analysis
    content = analysis["content"]
    assert "Visual Type" in content or "diagram" in content.lower()
    assert "Concept" in content or "glycolysis" in content.lower()


@pytest.mark.asyncio
async def test_pdf_visual_extraction(sample_pdf_with_diagram):
    """Verify PyMuPDF extracts embedded diagrams and creates visual knowledge units."""
    mat = Material(
        id="mat_pdf_test_1",
        title="Biochemistry Principles",
        material_type=MaterialType.TEXTBOOK,
        filename=sample_pdf_with_diagram.name,
        file_path=str(sample_pdf_with_diagram),
        subject="Biology",
    )

    units = await visual_extractor.extract_pdf_visuals(mat, sample_pdf_with_diagram)
    assert len(units) >= 1

    unit = units[0]
    assert unit.modality == ModalityType.VISUAL
    assert unit.source_tracking.page_number == 1
    assert unit.source_tracking.image_stored_path is not None
    assert unit.source_tracking.image_stored_path.startswith("/uploads/visuals/")

    # Citation pill format check: [🖼️ Figure X.X · PDF p.Y]
    citation = unit.source_tracking.citation_label
    assert "🖼️" in citation
    assert "PDF p.1" in citation

    # Ensure saved image file actually exists on disk
    rel_path = unit.source_tracking.image_stored_path.replace("/uploads/", "")
    full_disk_path = settings.UPLOAD_DIR / rel_path
    assert full_disk_path.exists()
    assert full_disk_path.stat().st_size > 0


@pytest.mark.asyncio
async def test_pptx_visual_extraction(sample_pptx_with_diagram):
    """Verify python-pptx extracts slide visuals and creates visual knowledge units."""
    mat = Material(
        id="mat_pptx_test_1",
        title="Bioenergetics Deck",
        material_type=MaterialType.SLIDE_DECK,
        filename=sample_pptx_with_diagram.name,
        file_path=str(sample_pptx_with_diagram),
        subject="Biology",
    )

    units = await visual_extractor.extract_slide_visuals(mat, sample_pptx_with_diagram)
    assert len(units) >= 1

    # Verify citation pill format: [🖼️ Slide X · ...]
    unit = units[0]
    assert unit.modality == ModalityType.VISUAL
    assert unit.source_tracking.image_stored_path.startswith("/uploads/visuals/")
    citation = unit.source_tracking.citation_label
    assert "🖼️" in citation
    assert "Slide" in citation


@pytest.mark.asyncio
async def test_end_to_end_ingestion_with_visuals(sample_pdf_with_diagram):
    """Verify IngestionService handles text + visual extraction without breaking."""
    mat = Material(
        id="mat_e2e_test_1",
        title="Cell Biology Textbook",
        material_type=MaterialType.TEXTBOOK,
        filename=sample_pdf_with_diagram.name,
        file_path=str(sample_pdf_with_diagram),
        subject="Biology",
    )
    repository.save_material(mat)

    await ingestion_service.process_material(
        task_id="task_e2e_vis_1",
        material=mat,
        file_path=sample_pdf_with_diagram,
    )

    # Check updated material status
    saved_mat = repository.get_material(mat.id)
    assert saved_mat is not None
    assert saved_mat.status == ProcessingStatus.INDEXED
    assert saved_mat.image_extraction_status == "complete"
    assert saved_mat.total_visuals_extracted >= 1
    assert saved_mat.total_units_extracted > saved_mat.total_visuals_extracted

    # Check units in repository
    all_units = repository.get_units_for_material(mat.id)
    text_units = [u for u in all_units if u.modality == ModalityType.TEXT]
    vis_units = [u for u in all_units if u.modality == ModalityType.VISUAL]

    assert len(text_units) >= 1
    assert len(vis_units) >= 1

    # Test static file serving endpoint for the extracted visual
    first_visual = vis_units[0]
    img_url = first_visual.source_tracking.image_stored_path
    resp = client.get(img_url)
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("image/")

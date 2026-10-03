import pytest
import io
from pathlib import Path
from pypdf import PdfWriter
from pptx import Presentation
from pptx.util import Inches, Pt
from app.models.material import Material, MaterialType
from app.models.knowledge_base import ModalityType
from app.services.parsers.pdf_parser import PDFTextbookParser
from app.services.parsers.video_parser import LectureVideoParser
from app.services.parsers.slide_parser import SlideDeckParser


@pytest.mark.asyncio
async def test_pdf_textbook_parser(tmp_path: Path):
    # Create a test PDF with 2 pages using pypdf
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.add_blank_page(width=200, height=200)

    pdf_path = tmp_path / "sample_textbook.pdf"
    with open(pdf_path, "wb") as f:
        writer.write(f)

    parser = PDFTextbookParser()
    assert parser.can_parse(pdf_path, MaterialType.TEXTBOOK)

    material = Material(
        title="Intro to Computer Science",
        material_type=MaterialType.TEXTBOOK,
        filename="sample_textbook.pdf",
        file_path=str(pdf_path),
        course_id="CS101",
        subject="Computer Science",
    )

    units = await parser.parse(material, pdf_path)
    # Even if blank, parser handles gracefully without error
    assert isinstance(units, list)


@pytest.mark.asyncio
async def test_lecture_video_parser_with_vtt(tmp_path: Path):
    vtt_content = """WEBVTT

00:00:10.000 --> 00:00:25.500
<v Professor Smith>Welcome everyone to Lecture 3 on Gradient Descent.

00:00:26.000 --> 00:00:45.000
<v Professor Smith>Today we will analyze convex optimization problems and convergence rates.
"""
    vtt_path = tmp_path / "lecture_03.vtt"
    vtt_path.write_text(vtt_content, encoding="utf-8")

    parser = LectureVideoParser()
    assert parser.can_parse(vtt_path, MaterialType.LECTURE_VIDEO)

    material = Material(
        title="Lecture 3: Optimization",
        material_type=MaterialType.LECTURE_VIDEO,
        filename="lecture_03.vtt",
        file_path=str(vtt_path),
        course_id="CS229",
    )

    units = await parser.parse(material, vtt_path)
    assert len(units) >= 1

    unit = units[0]
    assert unit.modality == ModalityType.SPEECH_TRANSCRIPT
    assert unit.source_tracking.start_time_seconds == 10.0
    assert unit.source_tracking.speaker_label == "Professor Smith"
    assert "Gradient Descent" in unit.content
    assert unit.source_tracking.start_timestamp == "00:10"
    assert "Lecture 3: Optimization" in unit.source_tracking.citation_label


@pytest.mark.asyncio
async def test_slide_deck_parser(tmp_path: Path):
    # Create a sample PPTX with python-pptx
    prs = Presentation()
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # Add title shape
    txBox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    tf = txBox.text_frame
    tf.text = "Convolutional Neural Networks"

    # Add speaker notes
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = "Remember to emphasize stride and padding here."

    pptx_path = tmp_path / "lecture_slides.pptx"
    prs.save(str(pptx_path))

    parser = SlideDeckParser()
    assert parser.can_parse(pptx_path, MaterialType.SLIDE_DECK)

    material = Material(
        title="Deep Learning Slides",
        material_type=MaterialType.SLIDE_DECK,
        filename="lecture_slides.pptx",
        file_path=str(pptx_path),
        course_id="CS231N",
    )

    units = await parser.parse(material, pptx_path)
    assert len(units) >= 2  # 1 for slide content, 1 for speaker notes

    # Check slide content unit
    slide_content_unit = next(u for u in units if not u.source_tracking.is_speaker_notes)
    assert slide_content_unit.modality == ModalityType.SLIDE_CONTENT
    assert slide_content_unit.source_tracking.slide_number == 1
    assert "Convolutional Neural Networks" in slide_content_unit.content

    # Check speaker notes unit
    notes_unit = next(u for u in units if u.source_tracking.is_speaker_notes)
    assert notes_unit.modality == ModalityType.SPEAKER_NOTES
    assert notes_unit.source_tracking.slide_number == 1
    assert "stride and padding" in notes_unit.content
    assert notes_unit.source_tracking.is_speaker_notes is True

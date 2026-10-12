import io
import re
import uuid
import logging
from pathlib import Path
from typing import List, Optional, Dict, Any, Tuple
from PIL import Image, ImageDraw, ImageFont

from app.core.config import settings
from app.models.material import Material, MaterialType
from app.models.knowledge_base import (
    KnowledgeUnit,
    ModalityType,
    SourceTrackingMetadata,
)
from app.services.vision_service import vision_service
from app.services.rag import embed_query

logger = logging.getLogger(__name__)

# Check pymupdf availability
try:
    import pymupdf
    PYMUPDF_AVAILABLE = True
except ImportError:
    try:
        import fitz as pymupdf
        PYMUPDF_AVAILABLE = True
    except ImportError:
        PYMUPDF_AVAILABLE = False

# Check python-pptx availability
try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False


class VisualExtractionPipeline:
    """
    Multimodal visual extraction engine for PDFs and Slide Decks.
    Detects embedded diagrams, charts, figures, and slide visuals,
    runs vision LLM semantic analysis, and constructs source-grounded
    visual KnowledgeUnits with image storage paths and citation pills.
    """

    def __init__(self, visual_dir: Optional[Path] = None):
        self.visual_dir = visual_dir or settings.VISUAL_DIR
        self.visual_dir.mkdir(parents=True, exist_ok=True)

    def _find_captions_on_page(self, page_text: str) -> List[str]:
        """Scans page text for standard academic figure/diagram caption labels."""
        pattern = r"(Figure\s+\d+(?:\.\d+)?[:\s\w\-,]{3,80}|Fig\.\s*\d+(?:\.\d+)?[:\s\w\-,]{3,80}|Chart\s+\d+[:\s\w\-,]{3,80}|Diagram\s+\d+[:\s\w\-,]{3,80})"
        matches = re.findall(pattern, page_text, re.IGNORECASE)
        return [m.strip() for m in matches if len(m.strip()) > 5]

    async def extract_pdf_visuals(
        self,
        material: Material,
        file_path: Path,
    ) -> List[KnowledgeUnit]:
        """
        Extracts embedded figures and diagrams from PDF textbooks using PyMuPDF (pymupdf/fitz).
        """
        visual_units: List[KnowledgeUnit] = []
        if not PYMUPDF_AVAILABLE:
            logger.warning("PyMuPDF (fitz) is not installed; PDF visual extraction skipped.")
            return visual_units

        logger.info(f"VisualExtraction: Inspecting PDF for diagrams/figures: {file_path.name}")
        try:
            doc = pymupdf.open(str(file_path))
            global_figure_idx = 0

            for page_idx in range(len(doc)):
                page_num = page_idx + 1
                page = doc[page_idx]
                page_text = page.get_text() or ""
                detected_captions = self._find_captions_on_page(page_text)

                image_list = page.get_images(full=True)
                if not image_list:
                    continue

                for img_local_idx, img_info in enumerate(image_list):
                    xref = img_info[0]
                    try:
                        base_image = doc.extract_image(xref)
                        if not base_image:
                            continue

                        image_bytes = base_image.get("image")
                        if not image_bytes or len(image_bytes) < 3000:
                            # Skip tiny icons, bullets, or glyphs
                            continue

                        width = base_image.get("width", 0)
                        height = base_image.get("height", 0)
                        if width < 120 or height < 120:
                            continue

                        global_figure_idx += 1
                        ext = base_image.get("ext", "png").lower()

                        # Determine caption
                        caption = None
                        if img_local_idx < len(detected_captions):
                            caption = detected_captions[img_local_idx]
                        elif detected_captions:
                            caption = detected_captions[0]
                        else:
                            caption = f"Figure {page_num}.{global_figure_idx}"

                        # Save image to storage with UUID
                        file_uuid = f"pdf_fig_{uuid.uuid4().hex[:12]}.{ext}"
                        disk_path = self.visual_dir / file_uuid
                        with open(disk_path, "wb") as f_out:
                            f_out.write(image_bytes)

                        stored_url = f"/uploads/visuals/{file_uuid}"

                        # Context for vision model
                        source_info = f"{material.title}, PDF Page {page_num}"
                        analysis = await vision_service.analyze_image(
                            image_bytes=image_bytes,
                            caption=caption,
                            context_text=page_text[:500],
                            source_info=source_info,
                            fallback_ext=ext,
                        )

                        visual_type = analysis.get("visual_type", "diagram")
                        desc = analysis.get("content", f"Diagram depicted in {caption}")

                        # Build citation label
                        # Format: [🖼️ Figure X.X · PDF p.Y]
                        fig_disp = caption if ("Figure" in caption or "Fig" in caption) else f"Figure {page_num}.{global_figure_idx}"
                        citation_label = f"[🖼️ {fig_disp} · PDF p.{page_num}]"

                        metadata = SourceTrackingMetadata(
                            material_id=material.id,
                            material_title=material.title,
                            material_type=MaterialType.TEXTBOOK,
                            chunk_index=global_figure_idx,
                            page_number=page_num,
                            page_range=f"p. {page_num}",
                            figure_index=global_figure_idx,
                            caption=caption,
                            visual_type=visual_type,
                            image_stored_path=stored_url,
                            citation_label=citation_label,
                        )

                        unit = KnowledgeUnit.create(
                            material_id=material.id,
                            content=desc,
                            modality=ModalityType.VISUAL,
                            source_tracking=metadata,
                            tags=[material.subject or "course", "visual", visual_type],
                            summary=f"{visual_type.title()} from page {page_num}: {caption}",
                        )
                        unit.embedding = embed_query(desc)
                        visual_units.append(unit)

                        logger.info(f"Extracted visual unit: {citation_label} (saved to {stored_url})")

                    except Exception as img_err:
                        logger.warning(f"Failed extracting image xref {xref} on page {page_num}: {img_err}")
                        continue

            doc.close()
            logger.info(f"Completed PDF visual extraction: found {len(visual_units)} visuals in {file_path.name}")
        except Exception as e:
            logger.error(f"Error in PDF visual extraction pipeline for {file_path.name}: {e}")

        return visual_units

    def _render_slide_composite(
        self,
        slide_title: str,
        body_points: List[str],
        embedded_img_bytes: Optional[bytes] = None,
    ) -> bytes:
        """
        Renders a clean high-resolution visual composite representing the slide
        using PIL. If embedded picture bytes exist, it embeds or composites it.
        """
        width, height = 1280, 720
        # Create 16:9 canvas
        canvas = Image.new("RGB", (width, height), color=(248, 250, 252))
        draw = ImageDraw.Draw(canvas)

        # Header bar
        draw.rectangle([(0, 0), (width, 80)], fill=(30, 41, 59))
        draw.text((40, 25), slide_title[:70] or "Lecture Slide", fill=(255, 255, 255))

        # If embedded image exists, paste it on the right half or center
        if embedded_img_bytes:
            try:
                emb = Image.open(io.BytesIO(embedded_img_bytes)).convert("RGB")
                emb.thumbnail((560, 560))
                canvas.paste(emb, (660, 100))
            except Exception:
                pass

        # Draw bullet items on the left side
        y_cursor = 120
        for pt in body_points[:8]:
            clean_pt = pt.replace("\n", " ")[:65]
            draw.text((50, y_cursor), f"* {clean_pt}", fill=(51, 65, 85))
            y_cursor += 50

        # Footer
        draw.rectangle([(0, height - 35), (width, height)], fill=(241, 245, 249))
        draw.text((40, height - 25), "AdaptFlow Multimodal Lecture Deck", fill=(148, 163, 184))

        out_buf = io.BytesIO()
        canvas.save(out_buf, format="PNG")
        return out_buf.getvalue()

    async def extract_slide_visuals(
        self,
        material: Material,
        file_path: Path,
    ) -> List[KnowledgeUnit]:
        """
        Extracts diagrams, figures, and rendered slide visuals from PowerPoint (.pptx).
        """
        visual_units: List[KnowledgeUnit] = []
        if not PPTX_AVAILABLE:
            logger.warning("python-pptx is not installed; Slide visual extraction skipped.")
            return visual_units

        logger.info(f"VisualExtraction: Inspecting slide deck for visuals: {file_path.name}")
        try:
            prs = Presentation(str(file_path))
            global_visual_idx = 0

            for slide_idx, slide in enumerate(prs.slides):
                slide_num = slide_idx + 1
                slide_title = "Slide Presentation"
                if slide.shapes.title and slide.shapes.title.text:
                    slide_title = slide.shapes.title.text.strip()

                # Gather slide bullet text & shape pictures
                slide_bullets: List[str] = []
                extracted_images: List[Tuple[bytes, str]] = []

                for shape in slide.shapes:
                    if shape == slide.shapes.title:
                        continue

                    # Check for picture shape
                    has_image = False
                    try:
                        if hasattr(shape, "shape_type") and shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                            has_image = True
                        elif hasattr(shape, "image"):
                            has_image = True
                    except Exception:
                        pass

                    if has_image:
                        try:
                            blob = shape.image.blob
                            ext = getattr(shape.image, "ext", "png")
                            if blob and len(blob) > 3000:
                                extracted_images.append((blob, ext))
                        except Exception:
                            pass

                    if shape.has_text_frame:
                        for p in shape.text_frame.paragraphs:
                            txt = p.text.strip()
                            if txt:
                                slide_bullets.append(txt)

                # Case 1: Extract individual embedded diagram pictures
                if extracted_images:
                    for img_idx, (img_bytes, img_ext) in enumerate(extracted_images):
                        global_visual_idx += 1
                        file_uuid = f"slide_{slide_num}_fig_{uuid.uuid4().hex[:10]}.{img_ext}"
                        disk_path = self.visual_dir / file_uuid
                        with open(disk_path, "wb") as f_out:
                            f_out.write(img_bytes)

                        stored_url = f"/uploads/visuals/{file_uuid}"
                        context_str = f"Slide #{slide_num}: {slide_title}\n" + "\n".join(slide_bullets[:5])
                        source_info = f"{material.title}, Slide #{slide_num}"

                        analysis = await vision_service.analyze_image(
                            image_bytes=img_bytes,
                            caption=f"{slide_title} (Diagram)",
                            context_text=context_str,
                            source_info=source_info,
                            fallback_ext=img_ext,
                        )

                        visual_type = analysis.get("visual_type", "diagram")
                        desc = analysis.get("content", f"Diagram on Slide {slide_num}")

                        # Pill format: [🖼️ Slide 4 · Diagram]
                        citation_label = f"[🖼️ Slide {slide_num} · {visual_type.capitalize()}]"

                        metadata = SourceTrackingMetadata(
                            material_id=material.id,
                            material_title=material.title,
                            material_type=MaterialType.SLIDE_DECK,
                            chunk_index=global_visual_idx,
                            slide_number=slide_num,
                            slide_title=slide_title,
                            visual_type=visual_type,
                            image_stored_path=stored_url,
                            citation_label=citation_label,
                        )

                        unit = KnowledgeUnit.create(
                            material_id=material.id,
                            content=desc,
                            modality=ModalityType.VISUAL,
                            source_tracking=metadata,
                            tags=[material.subject or "course", "slide_visual", visual_type],
                            summary=f"{visual_type.title()} on Slide {slide_num}: {slide_title}",
                        )
                        unit.embedding = embed_query(desc)
                        visual_units.append(unit)
                        logger.info(f"Extracted slide visual: {citation_label} (saved to {stored_url})")

                # Case 2: Render full slide visual representation
                else:
                    global_visual_idx += 1
                    composite_bytes = self._render_slide_composite(slide_title, slide_bullets)
                    file_uuid = f"slide_{slide_num}_rendered_{uuid.uuid4().hex[:10]}.png"
                    disk_path = self.visual_dir / file_uuid
                    with open(disk_path, "wb") as f_out:
                        f_out.write(composite_bytes)

                    stored_url = f"/uploads/visuals/{file_uuid}"
                    context_str = f"Slide #{slide_num}: {slide_title}\n" + "\n".join(slide_bullets[:6])
                    source_info = f"{material.title}, Slide #{slide_num}"

                    analysis = await vision_service.analyze_image(
                        image_bytes=composite_bytes,
                        caption=f"Slide {slide_num}: {slide_title}",
                        context_text=context_str,
                        source_info=source_info,
                        fallback_ext="png",
                    )

                    visual_type = analysis.get("visual_type", "diagram")
                    desc = analysis.get("content", f"Slide visual for {slide_title}")
                    citation_label = f"[🖼️ Slide {slide_num} · {visual_type.capitalize()}]"

                    metadata = SourceTrackingMetadata(
                        material_id=material.id,
                        material_title=material.title,
                        material_type=MaterialType.SLIDE_DECK,
                        chunk_index=global_visual_idx,
                        slide_number=slide_num,
                        slide_title=slide_title,
                        visual_type=visual_type,
                        image_stored_path=stored_url,
                        citation_label=citation_label,
                    )

                    unit = KnowledgeUnit.create(
                        material_id=material.id,
                        content=desc,
                        modality=ModalityType.VISUAL,
                        source_tracking=metadata,
                        tags=[material.subject or "course", "slide_visual", visual_type],
                        summary=f"Slide {slide_num} visual: {slide_title}",
                    )
                    unit.embedding = embed_query(desc)
                    visual_units.append(unit)

                    logger.info(f"Rendered slide visual unit: {citation_label} (saved to {stored_url})")

            logger.info(f"Completed slide visual extraction: found {len(visual_units)} visuals in {file_path.name}")
        except Exception as e:
            logger.error(f"Error in slide visual extraction pipeline for {file_path.name}: {e}")

        return visual_units


visual_extractor = VisualExtractionPipeline()
